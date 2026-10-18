#!/usr/bin/env python3
"""Map the structure of a thesis so the reviewer can plan a full read.

Usage:
    python thesis_map.py <thesis.pdf|.docx|.md|.txt> [--out DIR]

Writes to DIR (default: <system temp>/thesis_map/<file name>):
    thesis_text.txt   plain text of the whole document (read this, not the PDF)
    thesis_map.md     headings, word counts per chapter, figure/table counts,
                      approximate reference count, and RQ/objective lines found

The map is a navigation aid. Heading detection is heuristic; always confirm
against the text itself.
"""
import argparse
import re
import shutil
import subprocess
import sys
from pathlib import Path


# ---------- text extraction ----------

def extract_pdf(path: Path) -> list[tuple[str, int]]:
    """Return list of (line, level) with level 0 = body text."""
    text = ""
    if shutil.which("pdftotext"):
        r = subprocess.run(["pdftotext", "-layout", str(path), "-"],
                           capture_output=True, text=True)
        text = r.stdout
    if not text.strip():
        try:
            from pypdf import PdfReader
            text = "\n".join((p.extract_text() or "") for p in PdfReader(str(path)).pages)
        except Exception as e:  # pragma: no cover
            sys.exit(f"Could not extract PDF text: {e}")
    if len(text.split()) < 200:
        print("WARNING: very little text extracted — the PDF may be scanned "
              "images and need OCR.", file=sys.stderr)
    return [(ln.rstrip(), 0) for ln in text.splitlines()]


def extract_docx(path: Path) -> list[tuple[str, int]]:
    try:
        import docx
    except ImportError:
        r = subprocess.run(["pandoc", str(path), "-t", "markdown", "--wrap=none"],
                           capture_output=True, text=True)
        return extract_md_text(r.stdout)
    d = docx.Document(str(path))
    out = []
    for p in d.paragraphs:
        style = (p.style.name or "").lower() if p.style is not None else ""
        m = re.match(r"heading\s*(\d)", style)
        level = int(m.group(1)) if m else (1 if style == "title" else 0)
        out.append((p.text, level))
    # include table text so word counts aren't badly off
    for t in d.tables:
        for row in t.rows:
            out.append((" | ".join(c.text for c in row.cells), 0))
    return out


def extract_md_text(text: str) -> list[tuple[str, int]]:
    out = []
    for ln in text.splitlines():
        m = re.match(r"^(#{1,6})\s+(.*)", ln)
        out.append((m.group(2), len(m.group(1))) if m else (ln, 0))
    return out


# ---------- heuristic heading detection for plain text ----------

# Multilingual heading vocabulary. Covers English, Malay, Indonesian, French,
# Spanish, Portuguese, Italian, German, Dutch, Chinese and Arabic. Extend the
# lists if a thesis uses other terms; detection is case-insensitive.

def _alt(words):
    return "|".join(sorted(words, key=len, reverse=True))

CHAPTER_WORDS = ["chapter", "bab", "chapitre", "capítulo", "capitulo", "capitolo",
                 "kapitel", "hoofdstuk", "الفصل"]
NUMBER_WORDS = [r"\d{1,2}", r"[ivxlc]{1,6}",
                "one", "two", "three", "four", "five", "six", "seven", "eight", "nine", "ten",
                "satu", "dua", "tiga", "empat", "lima", "enam", "tujuh", "lapan", "delapan",
                "sembilan", "sepuluh", "un", "deux", "trois", "quatre", "cinq", "six", "sept",
                "uno", "dos", "tres", "cuatro", "cinco", "seis", "siete"]
# front/back matter: always top-level
FRONT_BACK = ["abstract", "abstrak", "résumé", "resumo", "resumen", "riassunto", "sommario",
              "zusammenfassung", "samenvatting", "摘要", "ملخص", "المستخلص",
              "acknowledgements?", "acknowledgments?", "penghargaan", "ucapan terima kasih",
              "kata pengantar", "remerciements", "agradecimientos", "danksagung", "شكر وتقدير",
              "references", "bibliography", "rujukan", "bibliografi", "daftar pustaka",
              "références", "bibliographie", "referencias", "bibliografía", "referências",
              "bibliografia", "literaturverzeichnis", "参考文献", "المراجع",
              "appendi(?:x|ces)", "lampiran", "annexes?", "anexos?", "allegati?", "anhang",
              "附录", "الملاحق",
              "table of contents", "contents", "isi kandungan", "kandungan", "daftar isi",
              "table des matières", "índice", "sumário", "indice", "inhaltsverzeichnis", "目录",
              "list of (?:tables|figures|abbreviations)", "senarai (?:jadual|rajah|singkatan)",
              "daftar (?:tabel|gambar|singkatan)"]
# section names that are a chapter when standing alone, a subsection inside a chapter
GENERIC = ["introduction", "pendahuluan", "pengenalan", "introducción", "introdução",
           "introduzione", "einleitung", "inleiding", "引言", "绪论", "مقدمة", "المقدمة",
           "conclusions?", "kesimpulan", "rumusan", "penutup", "conclusión", "conclusões",
           "conclusão", "conclusioni", "fazit", "schlussfolgerungen?", "conclusie", "结论",
           "الخاتمة",
           "literature review", "kajian literatur", "sorotan literatur", "tinjauan pustaka",
           "kajian pustaka", "revue de (?:la )?littérature", "revisión de (?:la )?literatura",
           "revisão de literatura", "文献综述",
           "methodology", "research methodology", "metodologi(?: kajian)?",
           "metode penelitian", "metodologi penelitian", "méthodologie", "metodología",
           "metodologia", "methodik", "methodologie", "研究方法", "منهجية البحث",
           "results", "findings", "keputusan", "dapatan(?: kajian)?", "hasil(?: penelitian| kajian)?",
           "résultats", "resultados", "risultati", "ergebnisse", "resultaten", "结果", "النتائج",
           "discussion", "perbincangan", "pembahasan", "discusión", "discussão", "discussione",
           "diskussion", "discussie", "讨论", "المناقشة",
           "summary", "ringkasan"]
FIG_WORDS = ["figure", "fig\\.", "rajah", "gambar", "figura", "abbildung", "abb\\.", "图", "شكل"]
TAB_WORDS = ["table", "jadual", "tabel", "tableau", "tabla", "tabella", "tabelle", "表", "جدول"]
REF_HEADINGS = ["references", "bibliography", "rujukan", "bibliografi", "daftar pustaka",
                "références", "bibliographie", "referencias", "bibliografía", "referências",
                "bibliografia", "literaturverzeichnis", "参考文献", "المراجع"]
APPX_HEADINGS = ["appendi(?:x|ces)", "lampiran", "annexes?", "anexos?", "allegati?", "anhang",
                 "附录", "الملاحق"]

CHAPTER_RE = re.compile(
    rf"^\s*(?:(?:{_alt(CHAPTER_WORDS)})\s+(?:{_alt(NUMBER_WORDS)})\b|第[一二三四五六七八九十\d]+章).{{0,90}}$",
    re.I)
SECTION_RE = re.compile(r"^\s*(\d{1,2})\.(\d{1,2})(\.\d{1,2})?\s+[^\W\d_][^\n]{2,90}$")
SPECIAL_RE = re.compile(rf"^\s*(?:\d{{1,2}}\.?\s+)?(?:{_alt(FRONT_BACK + GENERIC)})\s*[:.]?\s*\S{{0,15}}\s*$", re.I)
GENERIC_RE = re.compile(rf"^\s*(?:\d{{1,2}}\.?\s+)?(?:{_alt(GENERIC)})\b", re.I)
REF_RE = re.compile(rf"^\s*(?:{_alt(REF_HEADINGS)})\s*$", re.I | re.M)
APPX_RE = re.compile(rf"^\s*(?:{_alt(APPX_HEADINGS)})\b", re.I | re.M)


def detect_levels(lines: list[tuple[str, int]]) -> list[tuple[str, int]]:
    if any(lvl for _, lvl in lines):
        return lines  # styles already gave us headings
    out = []
    seen_chapter = False
    for ln, _ in lines:
        s = ln.strip()
        if CHAPTER_RE.match(s):
            seen_chapter = True
            out.append((s, 1))
        elif SPECIAL_RE.match(s):
            # "Conclusion"/"Kesimpulan" etc. inside a numbered chapter is usually a subsection
            out.append((s, 2 if (seen_chapter and GENERIC_RE.match(s)) else 1))
            if REF_RE.match(s) or APPX_RE.match(s):
                seen_chapter = False
        elif SECTION_RE.match(s) and not re.search(r"\.{4,}\s*\d+\s*$", s):
            out.append((s, 2 if not SECTION_RE.match(s).group(3) else 3))
        else:
            out.append((ln, 0))
    return out


def drop_toc(lines):
    """Remove table-of-contents style lines ('1.2 Background ....... 4')."""
    return [(ln, lvl if not re.search(r"(\.{3,}|\s{3,})\s*[0-9ivxlc]+\s*$", ln.strip()) else 0)
            for ln, lvl in lines]


# ---------- analysis ----------

CJK_RE = re.compile(r"[぀-ヿ㐀-䶿一-鿿가-힯]")


def words(s: str) -> int:
    """Word count for spaced scripts; CJK characters count one each (the usual convention)."""
    cjk = len(CJK_RE.findall(s))
    rest = CJK_RE.sub(" ", s)
    return cjk + len(re.findall(r"\w+(?:['’-]\w+)*", rest))


def build_map(lines):
    sections = []  # (title, level, wordcount)
    current = ["(Front matter before first heading)", 1, 0]
    for ln, lvl in lines:
        if lvl and lvl <= 3 and ln.strip():
            sections.append(tuple(current))
            current = [ln.strip(), lvl, 0]
        else:
            current[2] += words(ln)
    sections.append(tuple(current))
    return [s for s in sections if not (s[0].startswith("(Front") and s[2] == 0)]


def rollup_chapters(sections):
    chapters, cur = [], None
    for title, lvl, wc in sections:
        if lvl == 1:
            cur = [title, wc]
            chapters.append(cur)
        elif cur is not None:
            cur[1] += wc
        else:
            cur = ["(Front matter)", wc]
            chapters.append(cur)
    return chapters


def count_refs(full_text: str) -> int | None:
    m = list(REF_RE.finditer(full_text))
    if not m:
        return None
    tail = full_text[m[-1].end():]
    stop = APPX_RE.search(tail)
    if stop:
        tail = tail[:stop.start()]
    # entries usually start with "Surname, I." / "Surname I.," / "[n]" / "Surname, Name" (any script)
    L = r"[^\W\d_]"
    entries = re.findall(rf"^\s*(\[\d+\]|{L}[\w'’\-]*,\s*{L}\.|{L}[\w'’\-]+\s+{L}\.,|"
                         rf"{L}[\w'’\-]+(?:\s+{L}[\w'’\-]+)?,\s+{L}[\w'’\-]+.*\(\d{{4}}\w?\))",
                         tail, re.M)
    return len(entries)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("file")
    ap.add_argument("--out", default=None, help="output folder (default: a temp folder, never inside the skill)")
    a = ap.parse_args()
    path = Path(a.file)
    ext = path.suffix.lower()
    if ext == ".pdf":
        lines = extract_pdf(path)
    elif ext == ".docx":
        lines = extract_docx(path)
    elif ext in (".md", ".markdown"):
        lines = extract_md_text(path.read_text(errors="replace"))
    else:
        lines = [(ln, 0) for ln in path.read_text(errors="replace").splitlines()]

    lines = drop_toc(detect_levels(lines))
    full_text = "\n".join(ln for ln, _ in lines)

    import tempfile
    out = Path(a.out) if a.out else Path(tempfile.gettempdir()) / "thesis_map" / path.stem
    out.mkdir(parents=True, exist_ok=True)
    (out / "thesis_text.txt").write_text(full_text)

    sections = build_map(lines)
    chapters = rollup_chapters(sections)
    figs = len(set(re.findall(rf"(?:{_alt(FIG_WORDS)})\s*(\d+(?:\.\d+)?)", full_text, re.I)))
    tabs = len(set(re.findall(rf"(?<!\w)(?:{_alt(TAB_WORDS)})\s*(\d+(?:\.\d+)?)", full_text, re.I)))
    refs = count_refs(full_text)
    rq_terms = ["RQ\\s?\\d", "research questions?", "objectives?", "hypothes[ie]s", "H\\d",
                "soalan kajian", "persoalan kajian", "objektif", "hipotesis", "rumusan masalah",
                "tujuan penelitian", "pertanyaan penelitian", "questions? de recherche",
                "objectifs?", "preguntas? de investigación", "objetivos?", "forschungsfragen?",
                "研究问题", "研究目标", "أسئلة البحث", "أهداف البحث"]
    rq_re = re.compile(rf"(?<!\w)(?:{_alt(rq_terms)})(?!\w)", re.I)
    rq_lines = [ln.strip() for ln, _ in lines
                if rq_re.search(ln) and 4 < words(ln) < 80][:40]

    md = [f"# Thesis map: {path.name}", "",
          f"- Total words (approx.): {words(full_text):,}",
          f"- Distinct figures referenced: {figs}",
          f"- Distinct tables referenced: {tabs}",
          f"- Reference list entries (approx.): {refs if refs is not None else 'reference heading not found'}",
          "", "## Chapters (level-1 headings) with word counts", ""]
    md += [f"- {t} — {wc:,} words" for t, wc in chapters]
    md += ["", "## All headings", ""]
    md += [f"{'  ' * (lvl - 1)}- {t} ({wc:,} w)" for t, lvl, wc in sections]
    md += ["", "## Lines mentioning RQs / objectives / hypotheses (first 40)", ""]
    md += [f"- {l}" for l in rq_lines] or ["- none found"]
    md += ["", "_Heading detection is heuristic — confirm against thesis_text.txt._"]
    (out / "thesis_map.md").write_text("\n".join(md))
    print("\n".join(md))
    print(f"\nFull text: {out / 'thesis_text.txt'}")


if __name__ == "__main__":
    main()
